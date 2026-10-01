from __future__ import annotations

import html
import re

from io import BytesIO
from pathlib import Path

from docx import Document

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from fpdf import FPDF


def sanitize_text(text: str) -> str:

    text = text.replace(
        "\r\n",
        "\n",
    )

    text = text.replace(
        "\r",
        "\n",
    )

    replacements = {
        "“": '"',
        "”": '"',
        "‘": "'",
        "’": "'",
        "–": "-",
        "—": "-",
        "…": "...",
    }

    for old, new in replacements.items():

        text = text.replace(
            old,
            new,
        )

    return text.strip()


def _split_terms(
    terms: str,
) -> list[str]:

    return [
        part.strip()
        for part in terms.split(";")
        if part.strip()
    ]


def _set_cell_shading(
    cell,
    fill: str = "EDEDED",
) -> None:

    tc_pr = cell._tc.get_or_add_tcPr()

    shd = OxmlElement("w:shd")

    shd.set(
        qn("w:fill"),
        fill,
    )

    tc_pr.append(shd)


def format_docx(
    text: str,
    doc_type: str,
    terms: str = "",
    logo_path: Path | None = None,
) -> bytes:

    doc = Document()

    section = doc.sections[0]

    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = doc.styles["Normal"]

    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    # Logo
    if logo_path and logo_path.exists():

        paragraph = doc.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        paragraph.add_run().add_picture(
            str(logo_path),
            width=Inches(1.4),
        )

    # Title
    title = doc.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = title.add_run(
        doc_type.upper()
    )

    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(15)

    # Document body
    clean_text = sanitize_text(text)

    for line in clean_text.split("\n"):

        stripped = line.strip()

        if not stripped:

            doc.add_paragraph()

            continue

        paragraph = doc.add_paragraph()

        is_heading = (
            bool(
                re.match(
                    r"^(SECTION\s+)?\d+[.)]\s+",
                    stripped,
                    re.IGNORECASE,
                )
            )
            or (
                stripped.isupper()
                and len(stripped) < 100
            )
        )

        run = paragraph.add_run(
            stripped
        )

        run.font.name = "Times New Roman"

        if is_heading:
            run.bold = True

    # Terms table
    term_items = _split_terms(terms)

    if term_items:

        doc.add_heading(
            "Terms Summary",
            level=2,
        )

        table = doc.add_table(
            rows=1,
            cols=2,
        )

        table.alignment = (
            WD_TABLE_ALIGNMENT.CENTER
        )

        table.style = "Table Grid"

        header = table.rows[0].cells

        header[0].text = "#"
        header[1].text = "Term"

        for cell in header:

            _set_cell_shading(cell)

        for index, term in enumerate(
            term_items,
            start=1,
        ):

            cells = table.add_row().cells

            cells[0].text = str(index)
            cells[1].text = term

    # Footer
    footer = (
        section.footer.paragraphs[0]
    )

    footer.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer.text = (
        "LegalEase - AI-assisted draft. "
        "Review for your jurisdiction before use."
    )

    # Export
    output = BytesIO()

    doc.save(output)

    return output.getvalue()


class LegalEasePDF(FPDF):

    def __init__(
        self,
        title: str,
        logo_path: Path | None = None,
    ):

        super().__init__()

        self.doc_title = title
        self.logo_path = logo_path

        self.set_auto_page_break(
            auto=True,
            margin=18,
        )

    def header(self):

        if (
            self.logo_path
            and self.logo_path.exists()
        ):

            try:

                self.image(
                    str(self.logo_path),
                    x=92,
                    y=8,
                    w=26,
                )

                self.ln(18)

            except Exception:

                self.ln(4)

        self.set_font(
            "Times",
            "B",
            12,
        )

        self.cell(
            0,
            8,
            self.doc_title,
            align="C",
        )

        self.ln(10)

    def footer(self):

        self.set_y(-15)

        self.set_font(
            "Times",
            "I",
            8,
        )

        self.cell(
            0,
            10,
            (
                "LegalEase - AI-assisted draft. "
                "Review for your jurisdiction before use."
            ),
            align="C",
        )


def format_pdf(
    text: str,
    doc_type: str,
    logo_path: Path | None = None,
) -> bytes:

    pdf = LegalEasePDF(
        doc_type.upper(),
        logo_path,
    )

    pdf.set_title(doc_type)

    pdf.add_page()

    pdf.set_font(
        "Times",
        size=11,
    )

    clean_text = sanitize_text(text)

    for line in clean_text.split("\n"):

        stripped = line.strip()

        if not stripped:

            pdf.ln(4)

            continue

        is_heading = (
            bool(
                re.match(
                    r"^(SECTION\s+)?\d+[.)]\s+",
                    stripped,
                    re.IGNORECASE,
                )
            )
            or (
                stripped.isupper()
                and len(stripped) < 100
            )
        )

        pdf.set_font(
            "Times",
            "B" if is_heading else "",
            11,
        )

        pdf.multi_cell(
            0,
            6,
            stripped,
        )

        pdf.ln(1)

    return bytes(
        pdf.output()
    )


def format_html_preview(
    text: str,
) -> str:

    safe = html.escape(
        sanitize_text(text)
    )

    safe = re.sub(
        r"\n{2,}",
        "</p><p>",
        safe,
    )

    safe = safe.replace(
        "\n",
        "<br>",
    )

    return (
        '<div class="document-preview">'
        f"<p>{safe}</p>"
        "</div>"
    )