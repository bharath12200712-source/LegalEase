from pathlib import Path

from docx import Document

from backend.services.document_formatter import (
    format_docx,
    format_pdf,
    sanitize_text,
)


def test_sanitize_text():

    result = sanitize_text(
        "“Hello”—world…"
    )

    assert result == (
        '"Hello"-world...'
    )


def test_docx_export():

    data = format_docx(
        "1. Confidentiality\n"
        "The parties agree.",

        "NDA",

        (
            "Confidentiality must be maintained;"
            "No disclosure"
        ),
    )

    assert data.startswith(
        b"PK"
    )

    path = Path(
        "test_legalease.docx"
    )

    path.write_bytes(data)

    doc = Document(path)

    assert any(
        "NDA" in paragraph.text
        for paragraph in doc.paragraphs
    )

    assert len(doc.tables) == 1

    path.unlink()


def test_pdf_export():

    data = format_pdf(
        "1. Confidentiality\n"
        "The parties agree.",

        "NDA",
    )

    assert data.startswith(
        b"%PDF"
    )

    assert len(data) > 500